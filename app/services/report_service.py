import io
import calendar
from datetime import datetime
from typing import Optional, Dict, Any, List
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import or_, and_, desc, func

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import Cartridge, CartridgeStatus, HistoryLog, Branch, AppUser, Batch
from app.services.settings_service import SettingsService

STATUS_NAMES_RU = {
    CartridgeStatus.IN_USE: "В работе",
    CartridgeStatus.PENDING_VENDOR: "Ожидает заправщика",
    CartridgeStatus.AT_VENDOR: "На заправке",
    CartridgeStatus.READY_FOR_PICKUP: "Готов к выдаче"
}

MONTHS_RU = [
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"
]


class ReportService:
    @classmethod
    def get_report_data(
        cls,
        db: Session,
        current_user: AppUser,
        report_type: str = "all",
        branch_id: Optional[int] = None,
        year: Optional[int] = None,
        month: Optional[int] = None,
        date_from: Optional[str] = None,
        date_to: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Формирует структурированные данные отчета.
        Проверяет права доступа:
        - суперпользователь: любой филиал или все сразу
        - администратор/оператор с branch_id: строго только свой филиал
        - администратор/оператор без branch_id: любой филиал или все сразу
        """
        effective_branch_id = branch_id
        is_branch_locked = False
        if current_user:
            if current_user.role == "user":
                raise PermissionError("Доступ к формированию отчетов запрещен для вашей роли.")

            # Ограничение филиала по роли
            if current_user.role in ("admin", "operator") and current_user.branch_id:
                effective_branch_id = current_user.branch_id
                is_branch_locked = True

        branch_obj = None
        branch_name = "Все филиалы"
        if effective_branch_id:
            branch_obj = db.query(Branch).filter(Branch.id == effective_branch_id).first()
            if branch_obj:
                branch_name = branch_obj.name

        now = datetime.utcnow()
        dt_start = None
        dt_end = None
        period_title = "Все время"

        # 1. Определение временных рамок
        if report_type == "models":
            period_title = "Статистика по моделям картриджей"
        elif report_type == "year":
            y = year or now.year
            dt_start = datetime(y, 1, 1, 0, 0, 0)
            dt_end = datetime(y, 12, 31, 23, 59, 59)
            period_title = f"{y} год"
        elif report_type == "month":
            y = year or now.year
            m = month or now.month
            _, last_day = calendar.monthrange(y, m)
            dt_start = datetime(y, m, 1, 0, 0, 0)
            dt_end = datetime(y, m, last_day, 23, 59, 59)
            period_title = f"{MONTHS_RU[m - 1]} {y} г."
        elif report_type == "custom":
            try:
                dt_start = datetime.strptime(date_from, "%Y-%m-%d") if date_from else datetime(now.year, 1, 1)
            except Exception:
                dt_start = datetime(now.year, 1, 1)
            try:
                dt_end = datetime.strptime(date_to, "%Y-%m-%d").replace(hour=23, minute=59, second=59) if date_to else now
            except Exception:
                dt_end = now
            period_title = f"с {dt_start.strftime('%d.%m.%Y')} по {dt_end.strftime('%d.%m.%Y')}"
        else:
            report_type = "all"
            period_title = "Все картриджи за весь период"

        # 2. Выборка данных
        # Базовый запрос картриджей
        cart_query = db.query(Cartridge).options(
            joinedload(Cartridge.branch),
            joinedload(Cartridge.current_user)
        )
        if effective_branch_id:
            cart_query = cart_query.filter(Cartridge.branch_id == effective_branch_id)
        all_branch_cartridges = cart_query.all()
        cartridge_map = {c.id: c for c in all_branch_cartridges}

        if report_type == "models":
            # ОТЧЕТ ПО МОДЕЛЯМ КАРТРИДЖЕЙ
            models_map = {}
            for c in all_branch_cartridges:
                m_name = (c.model or "Не указана").strip()
                if m_name not in models_map:
                    models_map[m_name] = {
                        "model": m_name,
                        "total": 0,
                        "in_use": 0,
                        "pending_vendor": 0,
                        "at_vendor": 0,
                        "ready_for_pickup": 0,
                        "refill_count": 0,
                        "branch_name": branch_name
                    }
                entry = models_map[m_name]
                entry["total"] += 1
                if c.status == CartridgeStatus.IN_USE:
                    entry["in_use"] += 1
                elif c.status == CartridgeStatus.PENDING_VENDOR:
                    entry["pending_vendor"] += 1
                elif c.status == CartridgeStatus.AT_VENDOR:
                    entry["at_vendor"] += 1
                elif c.status == CartridgeStatus.READY_FOR_PICKUP:
                    entry["ready_for_pickup"] += 1

            if all_branch_cartridges:
                cart_ids = [c.id for c in all_branch_cartridges]
                refills_by_cart = dict(
                    db.query(HistoryLog.cartridge_id, func.count(HistoryLog.id))
                    .filter(
                        HistoryLog.cartridge_id.in_(cart_ids),
                        or_(HistoryLog.action.ilike("%возврат%"), HistoryLog.action.ilike("%приемка%"))
                    )
                    .group_by(HistoryLog.cartridge_id)
                    .all()
                )
                for c in all_branch_cartridges:
                    m_name = (c.model or "Не указана").strip()
                    if m_name in models_map:
                        models_map[m_name]["refill_count"] += refills_by_cart.get(c.id, 0)

            total_carts = len(all_branch_cartridges)
            models_list = list(models_map.values())
            models_list.sort(key=lambda x: x["total"], reverse=True)

            for m in models_list:
                m["percentage"] = round((m["total"] / total_carts * 100), 1) if total_carts > 0 else 0
                m["percentage_label"] = f"{m['percentage']}%"

            summary = {
                "total_models": len(models_list),
                "total_cartridges": total_carts,
                "in_use": sum(m["in_use"] for m in models_list),
                "pending_vendor": sum(m["pending_vendor"] for m in models_list),
                "at_vendor": sum(m["at_vendor"] for m in models_list),
                "ready_for_pickup": sum(m["ready_for_pickup"] for m in models_list),
                "total_refills": sum(m["refill_count"] for m in models_list),
                "top_model": models_list[0]["model"] if models_list else "—"
            }

            return {
                "report_type": "models",
                "period_title": period_title,
                "branch_name": branch_name,
                "branch_id": effective_branch_id,
                "is_branch_locked": is_branch_locked,
                "generated_at": now.strftime("%d.%m.%Y %H:%M"),
                "generated_by": current_user.full_name if current_user else "Система",
                "summary": summary,
                "items": models_list
            }

        elif report_type == "all":
            # ОБЩИЙ ОТЧЕТ: Полный реестр картриджей
            cartridges_data = []
            for c in all_branch_cartridges:
                user_display = "—"
                user_phone = "—"
                if c.current_user:
                    user_display = c.current_user.display_name
                    user_phone = c.current_user.phone or "—"
                elif c.current_user_id:
                    user_display = c.current_user_id

                refill_count = db.query(HistoryLog).filter(
                    HistoryLog.cartridge_id == c.id,
                    or_(
                        HistoryLog.action.ilike("%возврат%"),
                        HistoryLog.action.ilike("%приемка%")
                    )
                ).count()

                last_log = db.query(HistoryLog).filter(HistoryLog.cartridge_id == c.id).order_by(desc(HistoryLog.timestamp)).first()

                cartridges_data.append({
                    "id": c.id,
                    "marker_label": c.marker_label,
                    "qr_code": c.qr_code or "—",
                    "model": c.model,
                    "cabinet": c.cabinet,
                    "status": c.status.value if hasattr(c.status, "value") else str(c.status),
                    "status_label": STATUS_NAMES_RU.get(c.status, str(c.status)),
                    "branch_name": c.branch.name if c.branch else "Не указан",
                    "user_name": user_display,
                    "user_phone": user_phone,
                    "refill_count": refill_count,
                    "last_action": last_log.action if last_log else "Создание",
                    "last_action_date": last_log.timestamp.strftime("%d.%m.%Y %H:%M") if last_log else c.updated_at.strftime("%d.%m.%Y %H:%M") if c.updated_at else "—",
                    "notes": c.notes or ""
                })

            summary = {
                "total": len(cartridges_data),
                "in_use": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.IN_USE.value),
                "pending_vendor": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.PENDING_VENDOR.value),
                "at_vendor": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.AT_VENDOR.value),
                "ready_for_pickup": sum(1 for c in cartridges_data if c["status"] == CartridgeStatus.READY_FOR_PICKUP.value),
                "total_refills": sum(c["refill_count"] for c in cartridges_data)
            }

            return {
                "report_type": report_type,
                "period_title": period_title,
                "branch_name": branch_name,
                "branch_id": effective_branch_id,
                "is_branch_locked": is_branch_locked,
                "generated_at": now.strftime("%d.%m.%Y %H:%M"),
                "generated_by": current_user.full_name if current_user else "Система",
                "summary": summary,
                "items": cartridges_data
            }

        else:
            # ОТЧЕТ ЗА ПЕРИОД (год / месяц / интервал): Журнал движения и статистика
            log_query = db.query(HistoryLog).join(Cartridge).filter(
                HistoryLog.timestamp >= dt_start,
                HistoryLog.timestamp <= dt_end
            )
            if effective_branch_id:
                log_query = log_query.filter(Cartridge.branch_id == effective_branch_id)

            logs = log_query.order_by(desc(HistoryLog.timestamp)).all()

            # Подсчет показателей за период
            accepted_count = sum(1 for l in logs if "приемка" in l.action.lower())
            sent_vendor_count = sum(1 for l in logs if "передача" in l.action.lower() or "поставщик" in l.action.lower())
            returned_vendor_count = sum(1 for l in logs if "возврат" in l.action.lower())
            issued_count = sum(1 for l in logs if "выдача" in l.action.lower())
            wa_notifications = sum(1 for l in logs if "whatsapp" in l.action.lower())

            # Число актов за период
            batch_query = db.query(Batch).filter(
                Batch.created_at >= dt_start,
                Batch.created_at <= dt_end
            )
            if effective_branch_id:
                batch_query = batch_query.filter(Batch.branch_id == effective_branch_id)
            batches_count = batch_query.count()

            unique_cart_ids = set(l.cartridge_id for l in logs)

            # Формирование журнала операций
            log_items = []
            for l in logs:
                c = cartridge_map.get(l.cartridge_id) or l.cartridge
                b_name = "—"
                if c and c.branch:
                    b_name = c.branch.name
                elif branch_obj:
                    b_name = branch_obj.name

                log_items.append({
                    "id": l.id,
                    "timestamp": l.timestamp.strftime("%d.%m.%Y %H:%M"),
                    "cartridge_marker": c.marker_label if c else f"ID {l.cartridge_id}",
                    "cartridge_model": c.model if c else "—",
                    "branch_name": b_name,
                    "action": l.action,
                    "user_name": l.user_name or "—",
                    "details": l.details or ""
                })

            summary = {
                "total_operations": len(logs),
                "unique_cartridges": len(unique_cart_ids),
                "accepted_count": accepted_count,
                "sent_vendor_count": sent_vendor_count,
                "returned_vendor_count": returned_vendor_count,
                "issued_count": issued_count,
                "wa_notifications": wa_notifications,
                "batches_count": batches_count
            }

            return {
                "report_type": report_type,
                "period_title": period_title,
                "branch_name": branch_name,
                "branch_id": effective_branch_id,
                "is_branch_locked": is_branch_locked,
                "date_from": dt_start.strftime("%Y-%m-%d"),
                "date_to": dt_end.strftime("%Y-%m-%d"),
                "generated_at": now.strftime("%d.%m.%Y %H:%M"),
                "generated_by": current_user.full_name if current_user else "Система",
                "summary": summary,
                "items": log_items
            }

    @classmethod
    def generate_excel(cls, report: Dict[str, Any], org_name: str = "") -> io.BytesIO:
        """
        Создает стилизованный Excel-файл (.xlsx) с автоформатированием колонок и сводкой.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Отчет"
        ws.views.sheetView[0].showGridLines = True

        # Стили
        font_title = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        font_sub = Font(name="Calibri", size=10, italic=True, color="475569")
        font_meta = Font(name="Calibri", size=10, bold=True, color="1E293B")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=10, color="0F172A")
        font_kpi_num = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
        font_kpi_lbl = Font(name="Calibri", size=9, color="64748B")

        fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_kpi = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        border_thin = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin)

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        # 1. Шапка документа
        company = org_name or "Cartridge Tracker"
        if report["report_type"] == "models":
            report_title = "ОТЧЕТ ПО МОДЕЛЯМ КАРТРИДЖЕЙ"
        elif report["report_type"] == "all":
            report_title = "ОБЩИЙ ОТЧЕТ ПО КАРТРИДЖАМ"
        else:
            report_title = f"ОТЧЕТ ПО ОБОРОТУ КАРТРИДЖЕЙ ({report['period_title'].upper()})"

        ws.cell(row=1, column=1, value=company.upper()).font = font_sub
        ws.cell(row=2, column=1, value=report_title).font = font_title
        ws.cell(row=3, column=1, value=f"Филиал: {report['branch_name']}   |   Период: {report['period_title']}   |   Сформирован: {report['generated_at']} ({report['generated_by']})").font = font_meta

        current_row = 5

        # 2. Блок KPI сводки
        summary = report.get("summary", {})
        if report["report_type"] == "models":
            kpis = [
                ("Всего моделей", summary.get("total_models", 0)),
                ("Всего картриджей", summary.get("total_cartridges", 0)),
                ("В работе (у коллег)", summary.get("in_use", 0)),
                ("Ожидает заправщика", summary.get("pending_vendor", 0)),
                ("На заправке", summary.get("at_vendor", 0)),
                ("Готовы к выдаче", summary.get("ready_for_pickup", 0)),
                ("Всего заправок", summary.get("total_refills", 0)),
            ]
        elif report["report_type"] == "all":
            kpis = [
                ("Всего картриджей", summary.get("total", 0)),
                ("В работе (у коллег)", summary.get("in_use", 0)),
                ("Ожидает заправщика", summary.get("pending_vendor", 0)),
                ("На заправке", summary.get("at_vendor", 0)),
                ("Готовы к выдаче", summary.get("ready_for_pickup", 0)),
                ("Всего заправок", summary.get("total_refills", 0)),
            ]
        else:
            kpis = [
                ("Операций за период", summary.get("total_operations", 0)),
                ("Картриджей в обороте", summary.get("unique_cartridges", 0)),
                ("Принято в IT", summary.get("accepted_count", 0)),
                ("Отправлено поставщику", summary.get("sent_vendor_count", 0)),
                ("Возвращено / Заправлено", summary.get("returned_vendor_count", 0)),
                ("Выдано в работу", summary.get("issued_count", 0)),
            ]

        col_idx = 1
        for lbl, val in kpis:
            c_val = ws.cell(row=current_row, column=col_idx, value=val)
            c_val.font = font_kpi_num
            c_val.alignment = align_center
            c_val.fill = fill_kpi
            c_val.border = cell_border

            c_lbl = ws.cell(row=current_row + 1, column=col_idx, value=lbl)
            c_lbl.font = font_kpi_lbl
            c_lbl.alignment = align_center
            c_lbl.fill = fill_kpi
            c_lbl.border = cell_border
            col_idx += 1

        current_row += 3

        # 3. Основная таблица данных
        if report["report_type"] == "models":
            headers = [
                "№", "Модель картриджа", "Всего шт.", "Доля парка",
                "В работе", "Ожидает заправщика", "На заправке",
                "Готов к выдаче", "Всего заправок"
            ]
            fields = [
                "model", "total", "percentage_label",
                "in_use", "pending_vendor", "at_vendor",
                "ready_for_pickup", "refill_count"
            ]
        elif report["report_type"] == "all":
            headers = [
                "№", "Метка", "QR-код", "Модель", "Кабинет", "Статус",
                "Филиал", "Текущий владелец", "Телефон", "Кол-во заправок",
                "Последнее действие", "Дата обновления", "Примечания"
            ]
            fields = [
                "marker_label", "qr_code", "model", "cabinet", "status_label",
                "branch_name", "user_name", "user_phone", "refill_count",
                "last_action", "last_action_date", "notes"
            ]
        else:
            headers = [
                "№", "Дата / Время", "Картридж", "Модель", "Филиал",
                "Операция", "Исполнитель / Сотрудник", "Детали операции"
            ]
            fields = [
                "timestamp", "cartridge_marker", "cartridge_model", "branch_name",
                "action", "user_name", "details"
            ]

        # Запись заголовков таблицы
        for c_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=c_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = cell_border

        ws.row_dimensions[current_row].height = 25
        current_row += 1

        # Запись строк данных
        items = report.get("items", [])
        for idx, item in enumerate(items, 1):
            is_even = (idx % 2 == 0)
            row_fill = fill_zebra if is_even else None

            # Номер строки
            cell_num = ws.cell(row=current_row, column=1, value=idx)
            cell_num.font = font_data
            cell_num.alignment = align_center
            cell_num.border = cell_border
            if row_fill:
                cell_num.fill = row_fill

            for c_idx, fld in enumerate(fields, 2):
                val = item.get(fld, "")
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.font = font_data
                cell.border = cell_border
                if row_fill:
                    cell.fill = row_fill

                if fld in ("timestamp", "last_action_date", "refill_count", "cabinet", "status_label"):
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

            ws.row_dimensions[current_row].height = 20
            current_row += 1

        # 4. Автоматическая ширина колонок
        for col in ws.columns:
            max_len = 0
            for cell in col:
                # Пропускаем шапку и заголовок при расчете ширины
                if cell.row < 5:
                    continue
                v_str = str(cell.value or "")
                if len(v_str) > max_len:
                    max_len = len(v_str)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 50), 12)

        # Выгрузка в BytesIO
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
